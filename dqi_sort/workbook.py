"""Read past years' APS sorting workbooks (.xlsx) into the standard input and reference files.

    python -m dqi_sort.workbook extract XLSX OUT.csv --sheet NAME [--sheet NAME ...] [--virtual] [--no-sessions]
                                [--categories-from SHEET]
    python -m dqi_sort.workbook reference OUT.csv XLSX:SHEET [XLSX:SHEET ...]

For dry runs on earlier years. Each year's workbook is laid out differently, but every abstract table has a header
row with "ID" and "Title" (or, failing that, the standard APS column order); columns are found by their header
names (see HEADERS), and a sheet may hold several tables. Talks are deduplicated by ID (the first row wins; a later row only fills a missing session).

extract    writes the in-person or virtual input in the layout intake expects (the 2026 APS export columns).
           Invited talks listed without an abstract (the title is just the speaker's name) get a placeholder
           title and body naming their submission category, from the workbook's category list (a sheet with
           rows "17.01.05 | name"), so they can be carded and placed like the 2026 placeholders.
           --no-sessions leaves the session columns empty (use when the sheet is already sorted); --exclude drops the
           talks of another extracted file (a talk listed both in person and virtual stays virtual).
reference  writes the human final sort for scoring: talk_id, session_id, order, source (workbook and sheet).
           Sorting never reads it.
"""
import argparse, csv, os, re, sys

HEADERS = {
    'talk_id': ('id', 'abstract #'),
    'first': ('first name',),
    'last': ('last name',),
    'affiliation': ('affiliation',),
    'sc': ('sc', 'category'),
    'type': ('type', 'abstract type'),
    'title': ('title',),
    'body': ('body', 'abstract text'),
    'notes': ('submitter notes', 'author requests'),
    'tz': ('time zone/country', 'country', 'time zone'),
    'session': ('session abbreviation', 'session abrev', 'session title', 'session abbreviati'),
    'order': ('speaking order', 'order'),
    'sorter': ('sorter notes', 'sorter note'),
}
POSITIONAL = {'talk_id': 0, 'first': 1, 'last': 2, 'affiliation': 3, 'sc': 4, 'type': 5, 'title': 6, 'body': 7, 'notes': 8,
              'session': 9, 'order': 10, 'sorter': 11}
CAT = re.compile(r'^\d\d[a-z]?\.\d\d\.\d\d$')
IN_PERSON_COLS = ['ID', 'First Name', 'Last Name', 'Affiliation', 'SC', 'Type', 'Title', 'Body', 'Submitter Notes',
                  'Session Abbreviation', 'Speaking Order', 'Sorter Notes']
VIRTUAL_COLS = ['ID', 'First Name', 'Last Name', 'Affiliation', 'SC', 'Type', 'Title', 'Body', 'Submitter Notes',
                'Time Zone/Country', 'Session Abbreviation', 'Speaking Order', 'Sorter Notes']


def text(c):
    if c is None:
        return ''
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c).strip()


def header_map(cells):
    low = [c.lower() for c in cells]
    out = {}
    for field, names in HEADERS.items():
        for k, c in enumerate(low):
            if c in names and field not in out and k not in out.values():
                out[field] = k
                break
    return out if {'talk_id', 'title', 'type'} <= set(out) else None


def rows_of(path, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        sys.exit(f'{os.path.basename(path)} has no sheet {sheet!r}; sheets: {wb.sheetnames}')
    cols = None
    for r in wb[sheet].iter_rows(values_only=True):
        cells = [text(c) for c in r]
        h = header_map(cells)
        if h:
            cols = h
            continue
        if not cols:
            if len(cells) > 5 and cells[0].isdigit() and cells[5].lower() in ('oral', 'invited'):
                cols = POSITIONAL          # a table with no header row above it: the standard APS column order
            else:
                continue
        get = lambda f: cells[cols[f]] if f in cols and cols[f] < len(cells) else ''
        i = get('talk_id')
        if not i.isdigit() or get('type').lower() not in ('oral', 'invited'):
            continue
        yield {f: get(f) for f in HEADERS}


def categories(path, sheet):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for r in wb[sheet].iter_rows(values_only=True):
        cells = [text(c) for c in r]
        for k in range(len(cells) - 1):
            if CAT.match(cells[k]) and cells[k + 1] and not CAT.match(cells[k + 1]) and cells[k] not in out:
                out[cells[k]] = cells[k + 1]
    return out


def collect(path, sheets):
    talks = {}
    for sh in sheets:
        for r in rows_of(path, sh):
            t = talks.get(r['talk_id'])
            if t is None:
                talks[r['talk_id']] = dict(r, source=sh)
            elif not t['session'] and r['session']:
                t.update(session=r['session'], order=r['order'], source=sh)
    return talks


def extract(path, out, sheets, virtual=False, no_sessions=False, cat_sheet=None, exclude=None):
    talks = collect(path, sheets)
    if exclude:
        with open(exclude, encoding='utf-8', newline='') as f:
            drop = {r['ID'] for r in csv.DictReader(f)}
        talks = {i: t for i, t in talks.items() if i not in drop}
    cats = categories(path, cat_sheet) if cat_sheet else {}
    placeholders = 0
    rows = []
    for i in sorted(talks, key=lambda i: (len(i), i)):
        t = talks[i]
        title, body = t['title'], t['body']
        speaker = f"{t['first']} {t['last']}".strip()
        if t['type'].lower() == 'invited' and len(body) < 20 and (not title or title.lower() == speaker.lower() or len(title.split()) <= 3):
            name = cats.get(t['sc'], 'the submission category')
            title = f"Invited talk by {speaker} in {t['sc']}: {name}"
            body = (f'No title or abstract supplied for this invited talk by {speaker}; the topic is taken from its '
                    f'submission category, {t["sc"]} ({name}).')
            placeholders += 1
        row = {'ID': i, 'First Name': t['first'], 'Last Name': t['last'], 'Affiliation': t['affiliation'], 'SC': t['sc'],
               'Type': t['type'].capitalize(), 'Title': title, 'Body': body, 'Submitter Notes': t['notes'],
               'Time Zone/Country': t['tz'], 'Session Abbreviation': '' if no_sessions else t['session'],
               'Speaking Order': '' if no_sessions else t['order'], 'Sorter Notes': '' if no_sessions else t['sorter']}
        rows.append(row)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    with open(out, 'w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, VIRTUAL_COLS if virtual else IN_PERSON_COLS, extrasaction='ignore', lineterminator='\n')
        w.writeheader(); w.writerows(rows)
    print(f'{len(rows)} talks from {", ".join(sheets)} -> {out} ({placeholders} invited placeholders)')


def reference(out, specs):
    allt = {}
    for spec in specs:
        path, sheet = spec.rsplit(':', 1)
        for i, t in collect(path, [sheet]).items():
            if t['session'] and (i not in allt or not allt[i]['session']):
                allt[i] = dict(t, source=f'{os.path.basename(path)}:{sheet}')
    with open(out, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(['talk_id', 'session_id', 'order', 'source'])
        for i in sorted(allt, key=lambda i: (len(i), i)):
            w.writerow([i, allt[i]['session'], allt[i]['order'], allt[i]['source']])
    print(f'{len(allt)} sorted talks in {len({t["session"] for t in allt.values()})} sessions -> {out}')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest='cmd', required=True)
    e = sub.add_parser('extract')
    e.add_argument('xlsx'); e.add_argument('out'); e.add_argument('--sheet', action='append', required=True)
    e.add_argument('--virtual', action='store_true'); e.add_argument('--no-sessions', action='store_true')
    e.add_argument('--categories-from')
    e.add_argument('--exclude', help='a CSV written by extract whose IDs to leave out (e.g. the virtual talks)')
    r = sub.add_parser('reference')
    r.add_argument('out'); r.add_argument('specs', nargs='+')
    a = ap.parse_args()
    if a.cmd == 'extract':
        extract(a.xlsx, a.out, a.sheet, a.virtual, a.no_sessions, a.categories_from, a.exclude)
    else:
        reference(a.out, a.specs)


if __name__ == '__main__':
    main()
