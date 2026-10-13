import csv
import openpyxl
from dqi_sort import workbook as W


def make(path):
    wb = openpyxl.Workbook()
    home = wb.active; home.title = 'Home'
    home.append(['17.01.01', 'Made-up Qubits']); home.append(['17.01.02', 'Made-up Sensors'])
    ip = wb.create_sheet('In-person')
    ip.append(['Category summary', None, None])                       # no header: positional rows below
    ip.append([1000001, 'Ada', 'Lane', 'Univ A', '17.01.01', 'Oral', 'A made-up talk', 'Body text of a made-up talk.', '', '', '', ''])
    ip.append([1000002, 'Bo', 'Ng', 'Univ B', '17.01.02', 'Invited', 'Bo Ng', '', '', '', '', ''])
    srt = wb.create_sheet('Sorted')
    srt.append(['ID', 'First Name', 'Last Name', 'Affiliation', 'SC', 'Type', 'Title', 'Body', 'Submitter Notes', 'Session Abrev', 'Order'])
    srt.append([1000001, 'Ada', 'Lane', 'Univ A', '17.01.01', 'Oral', 'A made-up talk', 'Body.', '', 'SESS1', 2])
    srt.append([1000002, 'Bo', 'Ng', 'Univ B', '17.01.02', 'Invited', 'Bo Ng', '', '', 'SESS1', 1])
    wb.save(path)


def test_extract_with_placeholder_and_reference(tmp_path):
    x = tmp_path / 'w.xlsx'; make(x)
    W.extract(str(x), str(tmp_path / 'in.csv'), ['In-person'], no_sessions=True, cat_sheet='Home')
    rows = list(csv.DictReader(open(tmp_path / 'in.csv')))
    assert [r['ID'] for r in rows] == ['1000001', '1000002']
    assert rows[1]['Title'] == 'Invited talk by Bo Ng in 17.01.02: Made-up Sensors' and 'Bo Ng' in rows[1]['Body']
    W.reference(str(tmp_path / 'ref.csv'), [f'{x}:Sorted'])
    ref = list(csv.DictReader(open(tmp_path / 'ref.csv')))
    assert {(r['talk_id'], r['session_id'], r['order']) for r in ref} == {('1000001', 'SESS1', '2'), ('1000002', 'SESS1', '1')}
